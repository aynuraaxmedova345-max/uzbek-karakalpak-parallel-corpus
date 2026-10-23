#!/usr/bin/env python3
"""Loyiha fayllari manifesti: nom, hajm, SHA-256, Excel varaqlari/ustunlari/qatorlari,
PDF sahifalari, notebook kataklari. Fayllar faqat o'qiladi (Excel read_only rejimida).

  python build_manifest.py --root uzwordnet_audit --out uzwordnet_audit/output/manifest.json \
      [--extra PATH ...]

--extra: loyiha papkasidan tashqaridagi, lekin inventarga kiritilishi kerak bo'lgan fayllar.
"""
import argparse
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

SCAN_DIRS = ["input", "dictionaries"]


def sha256(p, buf=1 << 20):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        while chunk := f.read(buf):
            h.update(chunk)
    return h.hexdigest()


def xlsx_info(p):
    from openpyxl import load_workbook
    wb = load_workbook(p, read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        header, n = None, 0
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                header = [None if c is None else str(c) for c in row]
            if any(c is not None and str(c).strip() for c in row):
                n += 1
        sheets.append({"sheet": ws.title, "nonempty_rows_incl_first_row": n,
                       "data_rows_if_first_row_is_header": max(n - 1, 0), "first_row": header})
    wb.close()
    return {"sheets": sheets}


def pdf_info(p):
    from pypdf import PdfReader
    r = PdfReader(str(p))
    return {"pdf_pages": len(r.pages), "encrypted": r.is_encrypted}


def ipynb_info(p):
    nb = json.loads(Path(p).read_text(encoding="utf-8"))
    cells = nb.get("cells", [])
    return {"cells": len(cells), "code_cells": sum(c.get("cell_type") == "code" for c in cells)}


def describe(p, rel):
    p = Path(p)
    d = {"path": rel, "file_name": p.name, "size_bytes": p.stat().st_size, "sha256": sha256(p)}
    ext = p.suffix.lower()
    try:
        if ext in (".xlsx", ".xlsm"):
            d.update(xlsx_info(p))
        elif ext == ".pdf":
            d.update(pdf_info(p))
        elif ext == ".ipynb":
            d.update(ipynb_info(p))
    except Exception as e:  # o'qib bo'lmasa ham manifestda qoladi
        d["read_error"] = repr(e)
    return d


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--extra", nargs="*", default=[])
    ap.add_argument("--extra-label", nargs="*", default=[],
                    help="har bir --extra uchun manzil yorlig'i (masalan repo:branch:path)")
    a = ap.parse_args()
    root = Path(a.root)
    files = []
    for d in SCAN_DIRS:
        for p in sorted((root / d).rglob("*")):
            if p.is_file() and not p.name.startswith("."):
                files.append(describe(p, str(p.relative_to(root))))
    for i, p in enumerate(a.extra):
        label = a.extra_label[i] if i < len(a.extra_label) else str(p)
        files.append(describe(p, label))
    man = {"generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
           "scanned_dirs": SCAN_DIRS, "files": files}
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    Path(a.out).write_text(json.dumps(man, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"{len(files)} ta fayl -> {a.out}")


if __name__ == "__main__":
    main()
