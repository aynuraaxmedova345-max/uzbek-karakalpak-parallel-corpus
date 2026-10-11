#!/usr/bin/env python3
"""term_lexicons.ETYM (435 o'zak) ni O'zbek tilining izohli lug'ati bilan solishtirish.

Kirish: data/source/izohli_lugat_clean.xlsx ("etimologiya" ustuni, OCR qilingan).
Chiqish: output/kelib_chiqish_tekshiruv.xlsx (classify_terms.py uni terminlar_yakuniy.xlsx ga ham qo'shadi).
Lug'atga ham, ETYM ga ham hech narsa yozilmaydi.
"""
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook

import term_lexicons as L
from uztext import normalize

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data" / "source" / "izohli_lugat_clean.xlsx"
OUT = ROOT / "output" / "kelib_chiqish_tekshiruv.xlsx"
# Lug'at belgisi -> toifa. Faqat aniq o'qiladigan belgilar; OCR shovqini oldidagi [ { } 1 qoldiqlari olib tashlanadi.
MARKS = {"a.": "a.", "f.": "f.", "r.": "r."}
EUROPEAN = {"lot.", "yun.", "fr.", "nem.", "ingl.", "ital.", "isp.", "goll.", "ron.", "port.", "pol.", "lat.",
            "shved.", "ingl"}
MARK_RE = re.compile(r"^[\[\{\}1|]?\s*([a-zA-Z]+\.)")


def parse(et):
    """(belgi, toifa) yoki (None, sabab)."""
    if not et:
        return None, "belgisiz"
    s = et.strip()
    m = MARK_RE.match(s)
    if not m:
        return None, "o'qilmadi"
    tok = m.group(1).lower()
    if tok in MARKS and (m.group(1)[0].islower() or s[0] in "[{}1|"):
        nxt = re.match(r"^[\[\{\}1|]?\s*[a-z]+\.\s*-\s*([a-z]+\.)", s)   # a.-f. kabi
        if nxt and nxt.group(1) in MARKS:
            return f"{tok}-{nxt.group(1)}", "aralash"
        return tok, tok
    if tok in EUROPEAN:
        return tok, "yevropa"
    return None, "o'qilmadi"


def main():
    rows = list(load_workbook(SRC, read_only=True)["Izohli lugat"].iter_rows(values_only=True))
    h = rows[0]
    wi, ei = 0, h.index("etimologiya")
    entries = defaultdict(list)
    for r in rows[1:]:
        if r[wi]:
            entries[normalize(str(r[wi])).lower().strip()].append(r[ei])
    out = []
    for root, mine in sorted(L.ETYM.items()):
        head = next((c for c in (root, root + "moq", root + "mak") if c in entries), None)
        if head is None:
            out.append({"ozak": root, "mening_belgim": mine, "lugat_sozi": "", "lugat_belgisi": "",
                        "etimologiya_xom": "", "natija": "lug'atda topilmadi"})
            continue
        raws = [e for e in entries[head] if e]
        parsed = [parse(e) for e in raws] or [(None, "belgisiz")]
        mark, cat = next(((m, c) for m, c in parsed if m), parsed[0])
        if mark is None:
            res = ("belgi yo'q (o'zbekcha yoki OCR yo'qotgan)" if cat == "belgisiz" else "belgi o'qilmadi (OCR)")
            if mine == L.UZ and cat == "belgisiz":
                res = "mos (belgisiz = o'zbekcha)"
        elif mine == L.AP:
            res = (f"mos, aniqlashtirildi: {mark}" if cat in ("a.", "f.", "aralash") else f"farq: lug'atda {mark}")
        elif mine == L.RU:
            res = "mos" if cat in ("r.", "yevropa") else f"farq: lug'atda {mark}"
        else:
            res = f"farq: lug'atda {mark}"
        out.append({"ozak": root, "mening_belgim": mine, "lugat_sozi": head, "lugat_belgisi": mark or "",
                    "etimologiya_xom": " | ".join(str(e) for e in raws)[:200], "natija": res})
    order = lambda r: (0 if r["natija"].startswith("farq") else 1 if "topilmadi" in r["natija"]
                       else 2 if "o'qilmadi" in r["natija"] or "yo'q" in r["natija"] else 3, r["ozak"])
    out.sort(key=order)
    wb = Workbook(); ws = wb.active; ws.title = "kelib_chiqish_tekshiruv"
    cols = list(out[0]); ws.append(cols)
    for r in out:
        ws.append([r[c] for c in cols])
    wb.save(OUT)
    from collections import Counter
    c = Counter(re.sub(r":.*", "", r["natija"]) for r in out)
    print(len(out), dict(c))
    for r in out:
        if r["natija"].startswith("farq"):
            print(" ", r["ozak"], r["mening_belgim"], "->", r["natija"], "|", r["etimologiya_xom"][:60])


if __name__ == "__main__":
    main()
