#!/usr/bin/env python3
"""Avtomatik termin tasnifining aniqligi (precision) — qo'lda baholangan tanlanma bo'yicha.

output/terminlar_yakuniy.xlsx -> "tekshiruv_100" varag'idagi "mening_bahoim" ustuni to'ldirilgach:
    python scripts/precision.py [yo'l/terminlar_yakuniy.xlsx]

Qabul qilinadigan qiymatlar: to'g'ri / noto'g'ri (shuningdek: togri, notogri, +, -, 1, 0, ha, yo'q).
Bo'sh qatorlar hisobga olinmaydi. Natija: aniqlik, Wilson 95 % ishonch oralig'i, guruhlar kesimida,
hamda guruhlar ulushi bilan tortilgan (stratifikatsiyalangan) baho.
"""

import math
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
POS = {"to'g'ri", "toʻgʻri", "togri", "to‘g‘ri", "+", "1", "ha", "true"}
NEG = {"noto'g'ri", "notoʻgʻri", "notogri", "noto‘g‘ri", "-", "0", "yo'q", "yoʻq", "false"}
Z = 1.959964


def wilson(k, n, z=Z):
    if n == 0:
        return float("nan"), float("nan")
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return c - h, c + h


def main():
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "output" / "terminlar_yakuniy.xlsx"
    wb = load_workbook(path, read_only=True)
    rows = list(wb["tekshiruv_100"].iter_rows(values_only=True))
    head = rows[0]
    gi, vi = head.index("mavzuiy_guruh"), head.index("mening_bahoim")
    by = defaultdict(lambda: [0, 0])
    bad = []
    for r in rows[1:]:
        v = (str(r[vi]).strip().lower() if r[vi] is not None else "")
        if not v:
            continue
        if v in POS:
            by[r[gi]][0] += 1
        elif v not in NEG:
            bad.append(v)
            continue
        by[r[gi]][1] += 1
    if bad:
        print("Tanilmagan qiymatlar (hisobga olinmadi):", sorted(set(bad)))
    k = sum(a for a, _ in by.values())
    n = sum(b for _, b in by.values())
    if n == 0:
        sys.exit("Baholangan qator yo'q: 'mening_bahoim' ustunini to'ldiring.")
    lo, hi = wilson(k, n)
    print(f"Baholangan: {n} ta; to'g'ri: {k} ta")
    print(f"Aniqlik (precision): {k / n:.3f}  [Wilson 95 % ishonch oralig'i: {lo:.3f} – {hi:.3f}]")

    # guruh ulushlari yakuniy ro'yxatdan
    fin = list(wb["terminlar_yakuniy" if "terminlar_yakuniy" in wb.sheetnames else "terminlar_1000"]
               .iter_rows(values_only=True))
    fg = fin[0].index("mavzuiy_guruh")
    size = defaultdict(int)
    for r in fin[1:]:
        size[r[fg]] += 1
    total = sum(size.values())
    strat = sum(size[g] / total * (a / b) for g, (a, b) in by.items() if b)
    covered = sum(size[g] for g, (_, b) in by.items() if b) / total
    print(f"Guruh ulushlari bilan tortilgan aniqlik: {strat / covered:.3f} (qamrov: {covered:.0%} terminlar)")
    print("\nGuruhlar kesimida:")
    for g, (a, b) in sorted(by.items(), key=lambda x: -x[1][1]):
        l, h = wilson(a, b)
        print(f"  {g:36} {a:3}/{b:<3} = {a / b:.2f}  [{l:.2f} – {h:.2f}]")


if __name__ == "__main__":
    main()
