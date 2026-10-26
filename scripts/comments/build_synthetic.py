#!/usr/bin/env python3
"""Sun'iy (sintetik) izohlar datasetini yig'ish.

Kirish (data/synthetic/raw/):
  frazeologik_izohlar.txt, kochma_izohlar.txt
      qo'lda yozilgan izohlar: nomzod_id|izoh|sentiment|emotsiya|platforma
      sentiment: + - 0;  emotsiya: X G Q K H J N;  platforma: Y T I
  frazeologik_nomzodlar.csv, kochma_nomzodlar.csv
      leksikondan tanlangan birliklar va ularning ma'nolari (nomzod_id = cid)

Har bir izoh lexicon.py moslashtiruvchisi bilan tekshiriladi: unda aynan
mo'ljallangan ibora / ko'chma so'z topilishi shart. Tekshiruvdan o'tganlar
orasidan alifbo bo'ylab teng taqsimlangan N tasi tanlanadi.

  python scripts/comments/build_synthetic.py --n 1000
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lexicon import LexiconMatcher  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "data" / "synthetic" / "raw"
OUT = ROOT / "data" / "synthetic"

SENTIMENT = {"+": "ijobiy", "-": "salbiy", "0": "neytral"}
EMOTION = {"X": "xursandchilik", "G": "g'azab", "Q": "qayg'u", "K": "qo'rquv",
           "H": "hayrat", "J": "jirkanish", "N": "neytral"}
PLATFORM = {"Y": "YouTube", "T": "Telegram", "I": "Instagram"}

FIELDS = ["id", "kategoriya", "izoh", "birlik", "topilgan_qism", "manosi",
          "togri_manosi", "sentiment", "emotsiya", "platforma", "leksikon_manbasi"]

_APOS = "'‘’ʻʼ`´"


def original_span(text: str, norm_span: str) -> str:
    """Normallashtirilgan parchani asl matndan topib qaytaradi."""
    parts = []
    for ch in norm_span:
        if ch in _APOS:
            continue
        if ch.isspace() or ch == "-":
            parts.append(r"[\s\-–—]+")
        else:
            parts.append(re.escape(ch) + f"[{_APOS}]?")
    m = re.search("".join(parts), text, re.IGNORECASE)
    return m.group(0) if m else norm_span


def first_sentence(text: str, limit: int = 220) -> str:
    """Lug'at izohidan birinchi gapni (ta'rifni) ajratadi, misollarni tashlaydi."""
    text = re.sub(r"\s+", " ", text or "").strip()
    text = re.sub(r"^(?:s\.\s?t\.|esk\.|kt\.)\s*", "", text)  # uslub belgilari
    for m in re.finditer(r"[.;!?]\s*(?=…|\.\.|\s[A-ZА-ЯЎҚҒҲ«\"(-])", text):
        if m.start() >= 15:  # juda qisqa bo'lak (qisqartma, qavs) bo'lsa, davom etamiz
            text = text[: m.start() + 1]
            break
    return text if len(text) <= limit else text[:limit].rsplit(" ", 1)[0] + "…"


def read_rows(path: Path) -> list[list[str]]:
    rows = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.strip():
            parts = line.split("|")
            if len(parts) != 5:
                raise SystemExit(f"{path.name}: noto'g'ri qator: {line}")
            rows.append(parts)
    return rows


def read_candidates(path: Path) -> dict[int, dict]:
    with path.open(encoding="utf-8-sig", newline="") as fh:
        return {int(r["cid"]): r for r in csv.DictReader(fh)}


def spread(items: list, n: int) -> list:
    """Ro'yxatdan teng oraliqda n ta element tanlaydi (alifbo qamrovi uchun)."""
    if len(items) <= n:
        return items
    step = len(items) / n
    return [items[int(i * step)] for i in range(n)]


def build(kind: str, n: int) -> tuple[list[dict], int, int]:
    if kind == "frazeologik":
        rows = read_rows(RAW / "frazeologik_izohlar.txt")
        cands = read_candidates(RAW / "frazeologik_nomzodlar.csv")
        matcher = LexiconMatcher(use_kochma=False, gap=1)
    else:
        rows = read_rows(RAW / "kochma_izohlar.txt")
        cands = read_candidates(RAW / "kochma_nomzodlar.csv")
        matcher = LexiconMatcher(use_frazeologik=False, include_aux=True, gap=1)

    valid = []
    for cid, text, sent, emo, plat in rows:
        c = cands[int(cid)]
        if kind == "frazeologik":
            target = c["ibora"]
            hits = [m for m in matcher.find(text) if target in m.extra.get("ibora", "").split(" || ")]
            unit, meaning, literal = c["ibora_c"], first_sentence(c["m"]), ""
            source = f"frazeologik_lugat.csv, {c['manba_sahifa']}-bet"
        else:
            target = c["soz"]
            hits = [m for m in matcher.find(text) if target in m.item.split(" || ")]
            unit, meaning = c["soz"], first_sentence(c["m"])
            literal = first_sentence(c["togri_manoli_izoh"])
            source = f"kochma_leksikon.csv, id={c['id']}, {c['ma‘no_raqami']}-ma'no"
        if not hits:
            continue
        valid.append((int(cid), {
            "kategoriya": "frazeologik ibora" if kind == "frazeologik" else "ko'chma ma'noli so'z",
            "izoh": text, "birlik": unit, "topilgan_qism": original_span(text, hits[0].span),
            "manosi": meaning, "togri_manosi": literal,
            "sentiment": SENTIMENT[sent], "emotsiya": EMOTION[emo], "platforma": PLATFORM[plat],
            "leksikon_manbasi": source,
        }))
    valid.sort(key=lambda x: x[0])
    chosen = [r for _, r in spread(valid, n)]
    prefix = "F" if kind == "frazeologik" else "K"
    for i, r in enumerate(chosen, 1):
        r["id"] = f"{prefix}{i:04d}"
    return chosen, len(rows), len(valid)


def write_xlsx(path: Path, fraz: list[dict], koch: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from export_excel import (BODY_FONT, BORDER, FONT, HEADER_FILL, HEADER_FONT,
                              add_check_column, highlight, write_table)

    wb = Workbook()
    summary = wb.active
    summary.title = "Umumiy"
    lasts = {}
    for title, rows, kind in (("Frazeologik", fraz, "f"), ("Ko'chma", koch, "k")):
        ws = wb.create_sheet(title)
        headers = ["ID", "Izoh", "Birlik", "Topilgan parcha", "Ko'chma / frazeologik ma'nosi",
                   "To'g'ri ma'nosi", "Sentiment", "Emotsiya", "Platforma",
                   "Tekshiruv\n(ha / yo'q / shubhali)", "Leksikon manbasi"]
        data = [[r["id"], highlight(r["izoh"], [r["topilgan_qism"]]), r["birlik"], r["topilgan_qism"],
                 r["manosi"], r["togri_manosi"], r["sentiment"], r["emotsiya"], r["platforma"],
                 "", r["leksikon_manbasi"]] for r in rows]
        last = write_table(ws, headers, data,
                           widths=[8, 60, 26, 22, 50, 36 if kind == "k" else 6, 10, 14, 11, 14, 30],
                           wrap_cols=(2, 3, 4, 5, 6, 11), center_cols=(1, 7, 8, 9, 10),
                           input_cols=(10,))
        add_check_column(ws, 10, last)
        if kind == "f":
            ws.column_dimensions["F"].hidden = True
        lasts[title] = last

    ws = summary
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDE", (34, 16, 16, 16, 16)):
        ws.column_dimensions[col].width = w
    ws["A1"] = "Sun'iy izohlar dataseti: frazeologik iboralar va ko'chma ma'noli so'zlar"
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color="1F4E78")
    ws["A2"] = ("Izohlar leksikondagi ma'noga mos qilib qo'lda yozilgan va lexicon.py bilan "
                "tekshirilgan; har bir birlik datasetda bir marta uchraydi.")
    ws["A2"].font = Font(name=FONT, size=9, italic=True, color="595959")

    F, K = "Frazeologik", "'Ko''chma'"
    fl, kl = lasts["Frazeologik"], lasts["Ko'chma"]

    def header(row, labels):
        for c, label in enumerate(labels, 1):
            cell = ws.cell(row=row, column=c, value=label)
            cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, BORDER
            cell.alignment = Alignment(horizontal="left" if c == 1 else "center")

    def body(row, values):
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font, cell.border = BODY_FONT, BORDER
            if c > 1:
                cell.alignment = Alignment(horizontal="center")

    r = 4
    for title, col, values in (
        ("Sentiment", "G", ["ijobiy", "salbiy", "neytral"]),
        ("Emotsiya", "H", ["xursandchilik", "g'azab", "qayg'u", "qo'rquv", "hayrat", "jirkanish", "neytral"]),
        ("Platforma", "I", ["Telegram", "Instagram", "YouTube"]),
    ):
        ws.cell(row=r, column=1, value=title).font = Font(name=FONT, bold=True, color="1F4E78")
        r += 1
        header(r, [title, "Frazeologik", "Ko'chma", "Jami"])
        r += 1
        start = r
        for v in values:
            body(r, [v, f'=COUNTIF({F}!${col}$2:${col}${fl},A{r})',
                     f'=COUNTIF({K}!${col}$2:${col}${kl},A{r})', f"=B{r}+C{r}"])
            r += 1
        body(r, ["Jami", f"=SUM(B{start}:B{r-1})", f"=SUM(C{start}:C{r-1})", f"=SUM(D{start}:D{r-1})"])
        for c in range(1, 5):
            ws.cell(row=r, column=c).font = Font(name=FONT, size=10, bold=True)
        r += 2

    ws.cell(row=r, column=1, value="Tekshiruv holati").font = Font(name=FONT, bold=True, color="1F4E78")
    r += 1
    header(r, ["Belgi", "Frazeologik", "Ko'chma", "Jami"])
    for v in ("ha", "yo'q", "shubhali"):
        r += 1
        body(r, [v, f'=COUNTIF({F}!$J$2:$J${fl},A{r})', f'=COUNTIF({K}!$J$2:$J${kl},A{r})', f"=B{r}+C{r}"])
    r += 1
    body(r, ["tekshirilmagan", f"=COUNTBLANK({F}!$J$2:$J${fl})", f"=COUNTBLANK({K}!$J$2:$J${kl})", f"=B{r}+C{r}"])
    wb.save(path)


def write_csv(path: Path, rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(rows)


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--n", type=int, default=1000, help="har bir kategoriyadan nechta izoh")
    args = p.parse_args()

    fraz, f_total, f_valid = build("frazeologik", args.n)
    koch, k_total, k_valid = build("kochma", args.n)
    write_csv(OUT / "sintetik_frazeologik.csv", fraz)
    write_csv(OUT / "sintetik_kochma.csv", koch)
    write_csv(OUT / "sintetik_izohlar.csv", fraz + koch)
    write_xlsx(OUT / "sintetik_izohlar.xlsx", fraz, koch)
    print(f"frazeologik: yozilgan {f_total}, tekshiruvdan o'tgan {f_valid}, tanlangan {len(fraz)}")
    print(f"ko'chma:     yozilgan {k_total}, tekshiruvdan o'tgan {k_valid}, tanlangan {len(koch)}")


if __name__ == "__main__":
    main()
